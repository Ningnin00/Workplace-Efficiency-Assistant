# -*- coding: utf-8 -*-
"""
职场效率助手 - Web 应用后端

功能：
    - /api/weekly       生成周报
    - /api/ppt          生成 PPT 大纲（JSON）
    - /api/ppt/download 根据大纲生成 .pptx 文件下载

运行：
    python app.py

配置：
    设置环境变量 DEEPSEEK_API_KEY（或直接修改下方默认值）
"""

import os
import json
import re
import io
from flask import Flask, request, jsonify, send_file, send_from_directory
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ======================= 配置 =======================
# 在这里填入你的 DeepSeek API Key（或设置环境变量 DEEPSEEK_API_KEY）
DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "在这里填入你的_API_Key")
DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
MODEL = "deepseek-chat"

app = Flask(__name__, static_folder="static", static_url_path="")

# ======================= 提示词 =======================
WEEKLY_SYSTEM = """你是「职场效率助手」的周报专家，专门帮职场新人把零散的工作记录整理成规范、专业的周报。

【你的风格】
- 专业、简洁、有条理，用词准确不浮夸
- 像一位靠谱的资深同事在帮你把关
- 输出使用 Markdown 格式，层级清晰

【输出格式】
严格按以下结构输出周报：

## 本周工作
- 用「动词 + 具体事项 + 结果/进展」的句式，把零散记录归类整理

## 本周进展亮点
- 挑 1-2 个最有价值的成果，一句话说清楚价值

## 遇到的问题与解决
- 列出问题及解决办法，没有则写"无重大问题"

## 下周计划
- 3~5 条，用「目标 + 行动」的句式

【规则】
1. 输入太乱要先归类，不要照搬原文
2. 信息不足可合理推断，但不要编造具体数字
3. 语气积极专业，避免口语化废话
4. 最后附一句："需要我把它改成邮件语气发给 leader 吗？"
"""

PPT_SYSTEM = """你是「职场效率助手」的 PPT 策划专家，帮用户把一个主题扩展成结构清晰、逻辑完整的 PPT 大纲。

【你的风格】
- 逻辑清晰，金字塔结构：结论先行，层层展开
- 每页只讲一个重点，要点精炼（每条不超过 15 字）
- 适合职场汇报/方案展示场景

【输出格式】
必须只输出一段合法的 JSON，不要输出任何其他文字、不要 markdown 代码块标记，结构如下：
{
  "title": "PPT 主标题",
  "subtitle": "副标题",
  "author": "汇报人",
  "slides": [
    {"type": "cover", "title": "封面标题", "subtitle": "封面副标题", "author": "汇报人"},
    {"type": "agenda", "title": "目录", "points": ["第一部分", "第二部分"]},
    {"type": "content", "title": "页面标题", "points": ["要点1", "要点2"]},
    {"type": "end", "title": "谢谢观看", "subtitle": "结束语"}
  ]
}

【规则】
1. slides 至少包含：1 个 cover、1 个 agenda、若干 content、1 个 end
2. content 页建议 4~8 页，每页 points 3~5 条
3. 只输出 JSON，标题和要点都用中文
"""


# ======================= DeepSeek 调用 =======================
def call_deepseek(system_prompt, user_content):
    """调用 DeepSeek API，返回 (文本, 错误信息)"""
    if "在这里填入" in DEEPSEEK_API_KEY or not DEEPSEEK_API_KEY:
        return None, "未配置 DeepSeek API Key，请在 app.py 顶部填入你的 Key"

    headers = {
        "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_content},
        ],
        "temperature": 0.7,
        "stream": False,
    }
    try:
        resp = requests.post(DEEPSEEK_API_URL, headers=headers, json=payload, timeout=60)
        if resp.status_code != 200:
            return None, f"API 返回错误 {resp.status_code}: {resp.text[:200]}"
        data = resp.json()
        return data["choices"][0]["message"]["content"], None
    except Exception as e:
        return None, f"请求失败: {e}"


def extract_json(text):
    """从模型输出中提取 JSON 对象"""
    text = text.strip()
    # 去掉可能的 markdown 代码块
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.MULTILINE)
    try:
        return json.loads(text)
    except Exception:
        pass
    # 尝试截取第一个 { 到最后一个 }
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        try:
            return json.loads(text[start:end + 1])
        except Exception:
            pass
    return None


# ======================= PPT 生成 =======================
PRIMARY = RGBColor(0x2F, 0x54, 0xEB)
DARK = RGBColor(0x1F, 0x2A, 0x44)
LIGHT = RGBColor(0xF5, 0xF7, 0xFC)
ACCENT = RGBColor(0x00, 0xC8, 0x9A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height, text, size, color,
                 bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "微软雅黑"


def _add_bullets(slide, left, top, width, height, points, size=18, color=DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"\u25CF  {pt}"
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = "微软雅黑"
        p.space_after = Pt(14)


def _add_title_bar(slide, text):
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(0.55), Inches(0.12), Inches(0.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()
    _add_textbox(slide, Inches(0.9), Inches(0.5), Inches(11), Inches(0.7),
                 text, 30, DARK, bold=True)


def build_pptx_bytes(outline):
    """根据大纲生成 pptx 字节流"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for slide_data in outline.get("slides", []):
        stype = slide_data.get("type", "content")
        if stype == "cover":
            slide = prs.slides.add_slide(blank)
            _add_bg(slide, PRIMARY)
            _add_textbox(slide, Inches(1.5), Inches(2.2), Inches(10.3), Inches(1.5),
                         slide_data.get("title", ""), 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
            _add_textbox(slide, Inches(1.5), Inches(3.9), Inches(10.3), Inches(0.8),
                         slide_data.get("subtitle", ""), 20, WHITE, align=PP_ALIGN.CENTER)
            _add_textbox(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.6),
                         slide_data.get("author", ""), 16, WHITE, align=PP_ALIGN.CENTER)
        elif stype == "agenda":
            slide = prs.slides.add_slide(blank)
            _add_bg(slide, LIGHT)
            _add_title_bar(slide, slide_data.get("title", "目录"))
            _add_bullets(slide, Inches(1.2), Inches(1.6), Inches(11), Inches(5),
                         slide_data.get("points", []), size=22, color=PRIMARY)
        elif stype == "end":
            slide = prs.slides.add_slide(blank)
            _add_bg(slide, DARK)
            _add_textbox(slide, Inches(1.5), Inches(2.6), Inches(10.3), Inches(1.2),
                         slide_data.get("title", "谢谢观看"), 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
            _add_textbox(slide, Inches(1.5), Inches(4.0), Inches(10.3), Inches(0.8),
                         slide_data.get("subtitle", ""), 18, ACCENT, align=PP_ALIGN.CENTER)
        else:
            slide = prs.slides.add_slide(blank)
            _add_bg(slide, WHITE)
            _add_title_bar(slide, slide_data.get("title", ""))
            _add_bullets(slide, Inches(1.2), Inches(1.7), Inches(11), Inches(5),
                         slide_data.get("points", []), size=20)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ======================= 路由 =======================
@app.route("/")
def index():
    return send_from_directory("static", "index.html")


@app.route("/api/weekly", methods=["POST"])
def api_weekly():
    data = request.get_json(silent=True) or {}
    content = (data.get("content") or "").strip()
    if not content:
        return jsonify({"ok": False, "error": "请输入你的工作记录"})
    text, err = call_deepseek(WEEKLY_SYSTEM, content)
    if err:
        return jsonify({"ok": False, "error": err})
    return jsonify({"ok": True, "result": text})


@app.route("/api/ppt", methods=["POST"])
def api_ppt():
    data = request.get_json(silent=True) or {}
    topic = (data.get("topic") or "").strip()
    pages = data.get("pages") or 6
    if not topic:
        return jsonify({"ok": False, "error": "请输入 PPT 主题"})
    prompt = f"主题：{topic}，大约{pages}页左右"
    text, err = call_deepseek(PPT_SYSTEM, prompt)
    if err:
        return jsonify({"ok": False, "error": err})
    outline = extract_json(text)
    if not outline:
        return jsonify({"ok": False, "error": "AI 输出格式异常，请重试", "raw": text})
    return jsonify({"ok": True, "result": outline})


@app.route("/api/ppt/download", methods=["POST"])
def api_ppt_download():
    data = request.get_json(silent=True) or {}
    outline = data.get("outline")
    if not outline:
        return jsonify({"ok": False, "error": "缺少大纲数据"})
    try:
        buf = build_pptx_bytes(outline)
        filename = f"{outline.get('title', 'PPT')}.pptx"
        return send_file(buf, as_attachment=True, download_name=filename,
                         mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    except Exception as e:
        return jsonify({"ok": False, "error": f"生成 PPT 失败: {e}"})


if __name__ == "__main__":
    print("[OK] 职场效率助手已启动")
    print("[OK] 打开浏览器访问: http://127.0.0.1:5000")
    app.run(host="0.0.0.0", port=5000, debug=True)
